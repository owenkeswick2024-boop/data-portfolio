"""
生成演示用的“原始门店报表”：3 家门店 × 3 个月 = 9 个 Excel 文件。

故意模拟真实工作中常见的混乱情况：
- 表头上方有标题行和空行
- 不同门店的列名不一致（“日期”/“销售日期”、“商品”/“商品名称”……）
- 日期格式混杂（2026/7/3、2026-07-03、Excel 日期）
- 金额存成带千分位逗号的文本
- 商品名前后有多余空格
- 重复录入的行（同一单号出现两次）、空行、底部手工“合计”行
- 个别金额录错（与 数量×单价 不符）
"""
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

random.seed(42)
OUT = Path(__file__).parent / "input"
OUT.mkdir(exist_ok=True)

PRODUCTS = {  # 商品: 单价
    "无线鼠标": 89, "机械键盘": 299, "27寸显示器": 1299, "USB-C 扩展坞": 199,
    "降噪耳机": 599, "笔记本支架": 129, "移动硬盘 1TB": 399, "摄像头": 249,
}
STORES = {
    "深圳福田店": {"cols": ["单号", "日期", "商品", "数量", "单价", "金额"], "date_fmt": "slash", "prefix": "FT"},
    "深圳南山店": {"cols": ["订单号", "销售日期", "商品名称", "销量", "单价(元)", "销售额"], "date_fmt": "dash", "prefix": "NS"},
    "广州天河店": {"cols": ["流水号", "日期", "品名", "数量", "单价", "金额(元)"], "date_fmt": "excel", "prefix": "TH"},
}
MONTHS = [7, 8, 9]


def fmt_date(d, how):
    if how == "slash":
        return f"{d.year}/{d.month}/{d.day}"
    if how == "dash":
        return d.strftime("%Y-%m-%d")
    return d  # 真正的 Excel 日期


for store, cfg in STORES.items():
    for m in MONTHS:
        wb = Workbook()
        ws = wb.active
        ws.title = "销售流水"
        ws.append([f"{store} {m}月销售流水"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
        ws.append([])
        ws.append(cfg["cols"])

        rows = []
        start = date(2026, m, 1)
        days = (date(2026, m + 1, 1) - start).days
        for seq in range(1, random.randint(55, 80)):
            d = start + timedelta(days=random.randrange(days))
            p = random.choice(list(PRODUCTS))
            q = random.choices([1, 2, 3, 5], weights=[60, 25, 10, 5])[0]
            price = PRODUCTS[p]
            name = random.choice([p, f" {p}", f"{p} ", p])           # 多余空格
            amount = q * price
            amount_cell = f"{amount:,.2f}" if random.random() < 0.4 else amount  # 文本金额
            rows.append([f"{cfg['prefix']}{m:02d}{seq:04d}", fmt_date(d, cfg["date_fmt"]), name, q, price, amount_cell])

        rows.sort(key=lambda r: (str(r[1]), r[0]))
        typo = random.choice([r for r in rows if not isinstance(r[5], str)])  # 金额录错
        typo[5] = typo[3] * typo[4] + random.choice([-100, 10, 1000])
        for r in random.sample(rows, 3):                               # 重复录入
            rows.insert(random.randrange(len(rows)), list(r))
        for _ in range(2):                                             # 空行
            rows.insert(random.randrange(len(rows)), [None] * 6)

        for r in rows:
            ws.append(r)
            if cfg["date_fmt"] == "excel" and r[1] is not None:
                ws.cell(ws.max_row, 2).number_format = "yyyy/m/d"
        ws.append(["合计", None, None, None, None, "见上"])                    # 手工合计行

        for col, w in zip("ABCDEF", [12, 12, 16, 8, 10, 12]):
            ws.column_dimensions[col].width = w
        wb.save(OUT / f"{store}_2026年{m}月.xlsx")

print(f"已生成 {len(STORES) * len(MONTHS)} 个文件 → {OUT}")
